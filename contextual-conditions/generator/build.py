"""Build the Contextual Conditions proof-of-concept publication.

Reads the CDC CDSi release selected by StepIntoCDSI and the IVC-supplied
workbook, then writes static HTML, FHIR CodeSystem JSON, CSV, provenance and
validation reports under docs/contextual-conditions/.

Usage:
    py contextual-conditions/generator/build.py [--stepintocdsi PATH]

Requires openpyxl and PyYAML. Never downloads anything.
"""

import argparse
import collections
import csv
import hashlib
import html
import io
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl
import yaml

GENERATOR_VERSION = "0.1.0"

REPO = Path(__file__).resolve().parents[2]
CC_DIR = REPO / "contextual-conditions"
CONFIG_PATH = CC_DIR / "config" / "code-systems.json"
IVC_WORKBOOK = CC_DIR / "sources" / "ivc-supplied" / "Proposed Values.xlsx"
DOCS = REPO / "docs"
DEFAULT_STEPINTOCDSI = Path(r"C:/dev/cdsi/StepIntoCDSI")

VALUE_TYPE_MAP = {"Yes/No": "boolean", "Date": "date", "Number": "number"}
VALUE_TYPE_LABEL = {"boolean": "Yes/No", "date": "Date", "number": "Number"}

RELATED_SYSTEMS = {
    "SNOMED": ("relatedSnomedCt", "http://snomed.info/sct", "SNOMED CT"),
    "CVX": ("relatedCvx", "http://hl7.org/fhir/sid/cvx", "CVX"),
    "CDCPHINVS": ("relatedPhinVs", "urn:oid:2.16.840.1.114222.4.5.274", "PHIN VS"),
}

# Concept property definitions shared by both code systems. The URIs resolve
# to anchors on the root page, which documents each one.
PROPERTY_DEFS = {
    "expectedValueType": ("code", "Expected answer type of the assertion: boolean, date or number. Added by IVC; provisional vocabulary."),
    "indicationText": ("string", "CDSi indication text describing when the condition indicates vaccination."),
    "contraindicationText": ("string", "CDSi contraindication text describing when the condition contraindicates vaccination."),
    "clarifyingText": ("string", "CDSi clarifying text."),
    "relatedSnomedCt": ("Coding", "SNOMED CT concept related to this condition, as supplied by the source. Representative, not exhaustive, not equivalent."),
    "relatedCvx": ("Coding", "CVX vaccine code related to this condition, as supplied by the source. Representative, not exhaustive, not equivalent."),
    "relatedPhinVs": ("Coding", "PHIN VS concept related to this condition, as supplied by the source. Representative, not exhaustive, not equivalent."),
}


# ---------------------------------------------------------------- utilities

def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def esc(value):
    return html.escape("" if value is None else str(value))


def write_text(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")


def write_json(path, data):
    write_text(path, json.dumps(data, indent=2, ensure_ascii=False) + "\n")


def write_csv(path, header, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(header)
    w.writerows(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    # BOM so spreadsheet tools read the accented French correctly.
    path.write_bytes(("\ufeff" + buf.getvalue()).encode("utf-8"))


def cell_str(v):
    """Spreadsheet cell to the string a reader of the sheet sees."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


class Report:
    """Collects QA results: ('pass'|'info'|'warn'|'fail', check, detail)."""

    def __init__(self):
        self.items = []

    def add(self, level, check, detail=""):
        self.items.append((level, check, detail))

    def check(self, ok, check, detail="", fail_level="fail"):
        self.add("pass" if ok else fail_level, check, detail)

    def counts(self):
        return collections.Counter(level for level, _, _ in self.items)


# ---------------------------------------------------------------- sources

def load_cdsi(stepinto, report):
    sd = stepinto / "cdsi-reference" / "supporting-data"
    release = yaml.safe_load((sd / "current-version.yaml").read_text(encoding="utf-8"))["release_id"]
    rel_dir = sd / "versions" / str(release)
    manifest = yaml.safe_load((rel_dir / "manifest.yaml").read_text(encoding="utf-8"))
    report.add("info", "Release selected by StepIntoCDSI current-version.yaml", f"CDSi Supporting Data {release}")

    zip_path = rel_dir / "source" / manifest["source_filename"]
    zip_hash = sha256(zip_path)
    report.check(zip_hash == manifest["bundle_sha256"], "Source ZIP checksum matches manifest", f"{manifest['source_filename']} SHA-256 {zip_hash}")

    checked = {}
    for rel in ("xml/ScheduleSupportingData.xml", "spreadsheets/ScheduleSupportingData- Coded Observations-508.xlsx"):
        entry = next(f for f in manifest["files"] if f["path"] == rel)
        actual = sha256(rel_dir / "source" / rel)
        report.check(actual == entry["sha256"], f"Checksum matches manifest: {rel}", actual)
        checked[rel] = actual

    schedule = json.loads((rel_dir / "normalized" / "schedules" / "schedule.json").read_text(encoding="utf-8"))
    observations = schedule["data"]["observations"]["observation"]

    wb = openpyxl.load_workbook(rel_dir / "source" / "spreadsheets" / "ScheduleSupportingData- Coded Observations-508.xlsx", data_only=True)
    cond_rows = [r for r in wb["Conditions"].iter_rows(values_only=True) if any(c is not None for c in r)]
    sheet_header = [cell_str(c) for c in cond_rows[0]]
    sheet_rows = [[cell_str(c) for c in r] for r in cond_rows[1:]]
    overview = [cell_str(r[1]) for r in wb["Overview"].iter_rows(values_only=True) if len(r) > 1 and r[1]]
    history = []
    current = None
    for r in wb["Change History"].iter_rows(values_only=True):
        if r[0] == "Version":
            current = {"version": cell_str(r[1]), "published": cell_str(r[2]).replace("Publication Date:", "").strip(), "changes": []}
            history.append(current)
        elif current is not None and r[0] is None and r[1] is not None and r[0] != "Change":
            current["changes"].append({"area": cell_str(r[2]), "previous": cell_str(r[3]), "change": cell_str(r[4]), "reason": cell_str(r[5])})

    return {
        "release": str(release),
        "manifest": manifest,
        "zipSha256": zip_hash,
        "checkedFiles": checked,
        "observations": observations,
        "sheetHeader": sheet_header,
        "sheetRows": sheet_rows,
        "overview": overview,
        "history": history,
        "stepintoCommit": git_head(stepinto),
    }


def git_head(path):
    head = path / ".git" / "HEAD"
    try:
        ref = head.read_text().strip()
        if ref.startswith("ref:"):
            return (path / ".git" / ref[5:]).read_text().strip()
        return ref
    except OSError:
        return None


def load_ivc_workbook():
    wb = openpyxl.load_workbook(IVC_WORKBOOK, data_only=True)
    us = [r for r in wb["US"].iter_rows(values_only=True)]
    fr = [r for r in wb["FR"].iter_rows(values_only=True)]
    return us, fr


# ---------------------------------------------------------------- models

def coded_values(obs):
    cv = obs.get("codedValues") or {}
    cv = cv.get("codedValue") if isinstance(cv, dict) else None
    if cv is None:
        return []
    return cv if isinstance(cv, list) else [cv]


def build_cdsi_model(cdsi, us_rows, report):
    obs = cdsi["observations"]
    codes = [o["observationCode"] for o in obs]

    report.check(len(codes) == len(set(codes)), "Codes are unique", f"{len(codes)} codes")
    bad = [c for c in codes if not re.fullmatch(r"\d{3}", c)]
    report.check(not bad, "Codes are three-digit strings with leading zeroes preserved", ", ".join(bad))
    nums = sorted(int(c) for c in codes)
    gaps = [f"{n:03d}" for n in range(nums[0], nums[-1] + 1) if n not in set(nums)]
    report.add("info", "Unassigned numbers within the code range (preserved, never filled)", ", ".join(gaps) or "none")

    # Official spreadsheet vs normalized XML, field by field.
    sheet = {r[0]: r for r in cdsi["sheetRows"]}
    report.check(set(sheet) == set(codes), "Official spreadsheet and XML contain the same codes",
                 f"spreadsheet {len(sheet)}, XML {len(codes)}")
    mismatches = []
    mismatch_notes = collections.defaultdict(list)
    for o in obs:
        row = sheet.get(o["observationCode"])
        if not row:
            continue
        for idx, field in ((1, "observationTitle"), (2, "indicationText"), (3, "contraindicationText"), (4, "clarifyingText")):
            sv = "" if row[idx] == "n/a" else row[idx]
            xv = (o.get(field) or "").strip()
            if sv.strip() != xv:
                mismatches.append(f"{o['observationCode']} {field}: XML '{xv}' vs spreadsheet '{sv.strip()}'")
                mismatch_notes[o["observationCode"]].append(f"{field} differs from the official spreadsheet, which reads: '{sv.strip()}'. XML value shown")
    report.check(not mismatches, "Official spreadsheet text matches XML text (XML is used; differences flagged)", "; ".join(mismatches[:20]), fail_level="warn")

    # IVC-supplied workbook: expected value types plus a prior-snapshot comparison.
    ivc = {}
    int_codes = []
    for r in us_rows[1:]:
        if r[0] is None:
            continue
        raw = r[0]
        code = f"{int(raw):03d}" if isinstance(raw, (int, float)) else str(raw).strip()
        if isinstance(raw, (int, float)):
            int_codes.append(code)
        ivc[code] = r
    if int_codes:
        report.add("info", "IVC workbook stores some codes as numbers; re-padded to three digits for matching", ", ".join(int_codes))
    added = sorted(set(codes) - set(ivc))
    removed = sorted(set(ivc) - set(codes))
    report.add("info", "Codes added since the IVC-supplied workbook snapshot", ", ".join(added) or "none")
    report.check(not removed, "No codes removed since the IVC-supplied workbook snapshot", ", ".join(removed))
    retitled = [c for c in codes if c in ivc and str(ivc[c][1]).strip() != next(o for o in obs if o["observationCode"] == c)["observationTitle"].strip()]
    report.check(not retitled, "No titles changed since the IVC-supplied workbook snapshot", ", ".join(retitled), fail_level="warn")

    titles = collections.defaultdict(list)
    concepts = []
    for o in obs:
        code = o["observationCode"]
        titles[o["observationTitle"].strip()].append(code)
        src_type = ivc[code][9] if code in ivc else None
        flags = list(mismatch_notes.get(code, []))
        if src_type in VALUE_TYPE_MAP:
            vt, vt_src = VALUE_TYPE_MAP[src_type], "IVC workbook"
        else:
            vt, vt_src = "boolean", "IVC default (not in workbook)"
            flags.append("Expected value type not supplied; IVC default 'boolean' applied")
        related = []
        for cv in coded_values(o):
            prop, system, label = RELATED_SYSTEMS[cv["codeSystem"]]
            related.append({"property": prop, "system": system, "label": label, "code": cv["code"], "display": cv.get("text", "")})
        concepts.append({
            "code": code,
            "display": o["observationTitle"].strip(),
            "designations": [],
            "valueType": vt,
            "valueTypeSource": vt_src,
            "text": {k: (o.get(k) or "").strip() for k in ("indicationText", "contraindicationText", "clarifyingText")},
            "related": related,
            "flags": flags,
        })
    dups = {t: c for t, c in titles.items() if len(c) > 1}
    for t, c in dups.items():
        report.add("warn", "Display shared by more than one code (preserved; needs semantic review)", f"'{t}': {', '.join(c)}")
        for concept in concepts:
            if concept["code"] in c:
                concept["flags"].append(f"Display also used by {', '.join(x for x in c if x != concept['code'])}")
    defaulted = [c["code"] for c in concepts if c["valueTypeSource"].startswith("IVC default")]
    if defaulted:
        report.add("warn", "Expected value type defaulted by IVC (not in IVC workbook)", ", ".join(defaulted))
    vt_counts = collections.Counter(c["valueType"] for c in concepts)
    report.add("info", "Expected value type counts", ", ".join(f"{VALUE_TYPE_LABEL[k]} {v}" for k, v in sorted(vt_counts.items())))
    rel_counts = collections.Counter(r["label"] for c in concepts for r in c["related"])
    report.add("info", "Related code counts", ", ".join(f"{k} {v}" for k, v in sorted(rel_counts.items())))

    csv_header = cdsi["sheetHeader"] + ["IVC: expected value type", "IVC: expected value type source"]
    csv_rows = [row + [VALUE_TYPE_LABEL[c["valueType"]], c["valueTypeSource"]]
                for row, c in zip((sheet[c["code"]] for c in concepts), concepts)]
    return concepts, csv_header, csv_rows


FRENCH_CHARS = re.compile(r"[éèêàâçôûùîï]", re.I)


def build_syadem_model(fr_rows, report):
    header = [cell_str(c) for c in fr_rows[0]]
    if any("\xa0" in h for h in header):
        report.add("info", "Source column headers contain non-breaking spaces; normalized for output", ", ".join(repr(h) for h in header if h))
    body = fr_rows[1:]
    concept_rows = [r for r in body if r[0] is not None]
    other = [r for r in body if r[0] is None]
    summary = [r for r in other if any(c is not None for c in r)]
    report.add("info", "Non-concept spreadsheet rows excluded", f"{len(other)} rows ({len(other) - len(summary)} blank, {len(summary)} summary count rows)")
    report.add("info", "Empty spacer columns excluded", "columns D and F")

    codes = [str(r[0]).strip() for r in concept_rows]
    report.check(len(codes) == len(set(codes)), "Codes are unique", f"{len(codes)} codes")
    bad = [c for c in codes if not re.fullmatch(r"C-\d+", c)]
    report.check(not bad, "Codes use the C-<number> form", ", ".join(bad))
    types = collections.Counter(r[4] for r in concept_rows)
    unknown = [t for t in types if t not in VALUE_TYPE_MAP]
    report.check(not unknown, "Every Type is Yes/No, Date or Number", ", ".join(map(str, unknown)))
    expected = {r[4]: r[5] for r in summary if r[4]}
    report.check(all(types.get(t) == n for t, n in expected.items()), "Type counts match the workbook's own summary rows",
                 ", ".join(f"{t} {types.get(t)} (summary {n})" for t, n in expected.items()))

    concepts = []
    en_index = collections.defaultdict(list)
    fr_index = collections.defaultdict(list)
    ws_issues = []
    for r in concept_rows:
        code = str(r[0]).strip()
        en_raw, fr_raw = cell_str(r[1]), cell_str(r[2])
        en, frl = " ".join(en_raw.split()), " ".join(fr_raw.split())
        if en_raw != en or fr_raw != frl:
            ws_issues.append(code)
        en_index[en].append(code)
        fr_index[frl].append(code)
        flags = []
        if "\n" in en_raw:
            flags.append("English label contains a line break, possibly French and English combined")
        if FRENCH_CHARS.search(en):
            flags.append("English label appears to contain French text")
        elif en == frl:
            flags.append("English label identical to French (proper noun, shared term, or untranslated)")
        concepts.append({
            "code": code,
            "display": frl,
            "designations": [{"language": "en", "value": en}],
            "valueType": VALUE_TYPE_MAP[r[4]],
            "valueTypeSource": "Syadem",
            "text": {},
            "related": [],
            "flags": flags,
            "raw": [code if r[0] == code else cell_str(r[0]), en_raw, fr_raw, cell_str(r[4])],
        })
    if ws_issues:
        report.add("info", "Labels with extra, leading, trailing or non-breaking whitespace; collapsed in the CodeSystem and HTML, preserved in the CSV", f"{len(ws_issues)} codes: {', '.join(ws_issues[:15])}{' …' if len(ws_issues) > 15 else ''}")
    for label, index in (("English", en_index), ("French", fr_index)):
        for t, c in index.items():
            if len(c) > 1:
                report.add("warn", f"{label} label shared by more than one code (preserved; needs semantic review)", f"'{t}': {', '.join(c)}")
                for concept in concepts:
                    if concept["code"] in c:
                        concept["flags"].append(f"{label} label also used by {', '.join(x for x in c if x != concept['code'])}")
    for flag_name, check, level in (("appears to contain French", "English label appears to contain French text", "warn"),
                                    ("line break", "English label contains a line break (possibly combined French and English)", "warn"),
                                    ("identical to French", "English label identical to French (often proper nouns or shared terms)", "info")):
        hit = [c["code"] for c in concepts if any(flag_name in f for f in c["flags"])]
        if hit:
            report.add(level, check, f"{len(hit)} codes: {', '.join(hit)}")

    # Known semantic inconsistency called out in the planning note.
    for c in concepts:
        if c["code"] == "C-1029":
            c["flags"].append("English label does not match the French label or the Date type; source values preserved pending Syadem review")
            report.add("warn", "Label/type inconsistency flagged for Syadem review", "C-1029: English 'Intermediate living conditions', French 'Date du dernier antécédent de dengue', Type Date")

    vt_counts = collections.Counter(c["valueType"] for c in concepts)
    report.add("info", "Expected value type counts", ", ".join(f"{VALUE_TYPE_LABEL[k]} {v}" for k, v in sorted(vt_counts.items())))
    csv_header = ["Key", "English", "French", "Type"]
    return concepts, csv_header, [c["raw"] for c in concepts]


# ---------------------------------------------------------------- FHIR

def build_codesystem(cfg, cs, concepts, canonical, prop_base):
    used = {"expectedValueType"} | {k for c in concepts for k, v in c["text"].items() if v} | {r["property"] for c in concepts for r in c["related"]}
    props = [{"code": k, "uri": f"{prop_base}#{k}", "description": d, "type": t}
             for k, (t, d) in PROPERTY_DEFS.items() if k in used]
    out_concepts = []
    for c in concepts:
        entry = {"code": c["code"], "display": c["display"]}
        if c["designations"]:
            entry["designation"] = [{"language": d["language"], "value": d["value"]} for d in c["designations"]]
        p = [{"code": "expectedValueType", "valueCode": c["valueType"]}]
        for k in ("indicationText", "contraindicationText", "clarifyingText"):
            if c["text"].get(k):
                p.append({"code": k, "valueString": c["text"][k]})
        for r in c["related"]:
            coding = {"system": r["system"], "code": r["code"]}
            if r["display"]:
                coding["display"] = r["display"]
            p.append({"code": r["property"], "valueCoding": coding})
        entry["property"] = p
        out_concepts.append(entry)
    return {
        "resourceType": "CodeSystem",
        "id": cs["key"].replace("/", "-"),
        "language": cs["language"],
        "url": canonical,
        "version": cs["version"],
        "name": cs["name"],
        "title": cs["title"],
        "status": "draft",
        "experimental": True,
        "date": cfg["publicationDate"],
        "publisher": "Immunization Vocabularies Collaboration (IVC)",
        "contact": [{"name": "IVC", "telecom": [{"system": "email", "value": cfg["contactEmail"]}]}],
        "description": cs["description"] + "\n\n**Proof of concept.** This resource is not a final or normative release. Its structure and metadata are still under review. Source identifiers are permanent.",
        "copyright": cs["copyright"],
        "caseSensitive": True,
        "compositional": False,
        "versionNeeded": False,
        "content": "complete",
        "count": len(out_concepts),
        "property": props,
        "concept": out_concepts,
    }


# ---------------------------------------------------------------- HTML

NAV = [("home.html", "Home"), ("nuva.html", "NUVA"), ("resources.html", "Resources"), ("meetings.html", "Meetings"), ("about.html", "About")]


def page(title, description, depth, body, eyebrow, h1, lead, crumbs=()):
    up = "../" * depth
    nav = "".join(f'<li><a href="{up}{href}">{label}</a></li>' for href, label in NAV)
    crumb_html = ""
    if crumbs:
        parts = [f'<a href="{esc(h)}">{esc(t)}</a>' if h else f'<span aria-current="page">{esc(t)}</span>' for t, h in crumbs]
        crumb_html = f'<nav class="cc-crumbs" aria-label="Breadcrumb">{" / ".join(parts)}</nav>'
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><meta name="robots" content="noindex"><meta name="description" content="{esc(description)}"><title>{esc(title)} | Immunization Vocabularies Collaboration</title><link rel="icon" type="image/svg+xml" href="{up}favicon.svg"><link rel="stylesheet" href="{up}site.css"><link rel="stylesheet" href="{up}contextual-conditions/cc.css"></head>
<body><a class="skip-link" href="#main">Skip to main content</a>
<header class="site-header"><div class="wrap header-inner"><a class="brand" href="{up}home.html" aria-label="IVC home"><img src="{up}assets/ivc-logo.png" alt="Immunization Vocabularies Collaboration"></a><button class="menu-toggle" type="button" aria-expanded="false" aria-controls="primary-nav">Menu</button><nav class="primary-nav" id="primary-nav" aria-label="Primary navigation"><ul class="nav-list">{nav}<li><a class="nav-cta" href="{up}contact.html">Contact</a></li></ul></nav></div></header>
<div class="cc-poc" role="note"><div class="wrap"><strong>Proof of concept: not a final or normative release.</strong> This page shows how IVC might publish contextual condition code systems. Governance, URLs, metadata and structure are still under discussion. Source codes are permanent, but everything else here may change.</div></div>
<main id="main"><section class="page-hero cc-hero"><div class="wrap">{crumb_html}<p class="eyebrow">{esc(eyebrow)}</p><h1>{esc(h1)}</h1><p class="lead">{lead}</p></div></section>
{body}</main>
<footer class="site-footer"><div class="wrap"><div class="footer-grid"><div class="footer-brand"><img src="{up}assets/ivc-logo.png" alt=""><p>Supporting the people who make immunization data understandable.</p></div><div><h2>Explore</h2><ul><li><a href="{up}nuva.html">NUVA</a></li><li><a href="{up}resources.html">Resources</a></li><li><a href="{up}meetings.html">Meetings</a></li></ul></div><div><h2>IVC</h2><ul><li><a href="{up}about.html">About</a></li><li><a href="{up}contact.html">Contact</a></li><li><a href="mailto:info@ivci.org">info@ivci.org</a></li></ul></div></div><div class="footer-note">Immunization Vocabularies Collaboration · Contextual Conditions proof of concept · generated by build.py {GENERATOR_VERSION}</div></div></footer><script src="{up}site.js"></script></body></html>
"""


def kv_table(rows):
    return '<dl class="cc-kv">' + "".join(f"<div><dt>{esc(k)}</dt><dd>{v}</dd></div>" for k, v in rows) + "</dl>"


def roles_table(roles):
    labels = [("codeSystemAuthority", "Code System Authority", "Defines what each concept means."),
              ("publishingAuthority", "Publishing Authority", "Publishes or distributes the vocabulary."),
              ("namespaceAuthority", "Namespace Authority", "Maintains the canonical URL and unique identifiers."),
              ("cdsAuthority", "CDS Authority", "Decides how concepts affect recommendation logic.")]
    rows = "".join(f'<tr><th scope="row">{esc(l)}<span class="cc-sub">{esc(d)}</span></th><td>{esc(roles[k])}</td></tr>' for k, l, d in labels)
    return f'<div class="cc-scroll"><table class="cc-table cc-roles"><thead><tr><th scope="col">Role</th><th scope="col">Held by</th></tr></thead><tbody>{rows}</tbody></table></div>'


def tag(kind, text):
    return f'<span class="cc-tag cc-tag-{kind}">{esc(text)}</span>'


def hl7_examples(cs, canonical, concepts):
    by_code = {c["code"]: c for c in concepts}
    lines = []
    for i, ex in enumerate(cs["hl7v2Examples"], 1):
        c = by_code[ex["code"]]
        lines.append(f'OBX|{i}|{ex["valueType"]}|{c["code"]}^{c["display"]}^{canonical}|{i}|{ex["value"]}||||||F')
    first = by_code[cs["hl7v2Examples"][0]["code"]]
    coding = json.dumps({"system": canonical, "version": cs["version"], "code": first["code"], "display": first["display"]}, indent=2, ensure_ascii=False)
    return f"""<h3>HL7 v2 (OBX in VXU or QBP)</h3>
<pre class="cc-code"><code>{esc(chr(10).join(lines))}</code></pre>
<p class="meta">These examples follow the proposal: OBX-3 carries the code, its display and this canonical URL. A Yes/No assertion uses value type ID with HL7 table 0136. A date-valued assertion uses DT. The patient data is illustrative.</p>
<h3>FHIR Coding</h3>
<pre class="cc-code"><code>{esc(coding)}</code></pre>"""


FILTER_JS = """<script>
(function () {
  var input = document.getElementById('cc-filter');
  var type = document.getElementById('cc-type');
  var flagged = document.getElementById('cc-flagged');
  var rows = Array.prototype.slice.call(document.querySelectorAll('#cc-concepts tbody tr'));
  var count = document.getElementById('cc-count');
  function apply() {
    var q = input.value.trim().toLowerCase();
    var t = type.value;
    var f = flagged.checked;
    var shown = 0;
    rows.forEach(function (r) {
      var ok = (!q || r.textContent.toLowerCase().indexOf(q) !== -1) &&
               (!t || r.getAttribute('data-type') === t) &&
               (!f || r.hasAttribute('data-flagged'));
      r.hidden = !ok;
      if (ok) shown++;
    });
    count.textContent = shown + ' of ' + rows.length + ' concepts shown';
  }
  input.addEventListener('input', apply);
  type.addEventListener('change', apply);
  flagged.addEventListener('change', apply);
  apply();
})();
</script>"""


def concept_table(cs, concepts):
    is_cdsi = cs["source"] == "cdsi"
    if is_cdsi:
        head = '<th scope="col">Code</th><th scope="col">Display</th><th scope="col">Answer <span class="cc-mini">IVC</span></th><th scope="col">Indication</th><th scope="col">Contraindication</th><th scope="col">Clarifying text</th><th scope="col">Related codes</th>'
    else:
        head = '<th scope="col">Code</th><th scope="col">Français (display)</th><th scope="col">English</th><th scope="col">Answer</th>'
    rows = []
    for c in concepts:
        flags = "".join(f'<li>{esc(f)}</li>' for f in c["flags"])
        flag_html = f'<ul class="cc-flags">{flags}</ul>' if flags else ""
        vt = esc(VALUE_TYPE_LABEL[c["valueType"]])
        if c["valueTypeSource"].startswith("IVC default"):
            vt += f' <span class="cc-mini" title="{esc(c["valueTypeSource"])}">IVC default</span>'
        attrs = f' id="c-{esc(c["code"])}" data-type="{c["valueType"]}"' + (" data-flagged" if c["flags"] else "")
        code_cell = f'<td class="cc-codecell"><a href="#c-{esc(c["code"])}">{esc(c["code"])}</a></td>'
        if is_cdsi:
            rel = "".join(f'<li><span class="cc-sys">{esc(r["label"])}</span> {esc(r["code"])}{(" " + esc(r["display"])) if r["display"] else ""}</li>' for r in c["related"])
            rows.append(f'<tr{attrs}>{code_cell}<td class="cc-disp">{esc(c["display"])}{flag_html}</td><td>{vt}</td>'
                        f'<td>{esc(c["text"]["indicationText"])}</td><td>{esc(c["text"]["contraindicationText"])}</td><td>{esc(c["text"]["clarifyingText"])}</td>'
                        f'<td>{"<ul class=cc-rel>" + rel + "</ul>" if rel else ""}</td></tr>')
        else:
            rows.append(f'<tr{attrs}>{code_cell}<td class="cc-disp" lang="fr">{esc(c["display"])}{flag_html}</td><td lang="en">{esc(c["designations"][0]["value"])}</td><td>{vt}</td></tr>')
    types = sorted({c["valueType"] for c in concepts})
    opts = "".join(f'<option value="{t}">{VALUE_TYPE_LABEL[t]}</option>' for t in types)
    n_flag = sum(1 for c in concepts if c["flags"])
    return f"""<div class="cc-filters">
<label>Search <input id="cc-filter" type="search" placeholder="Code or text" autocomplete="off"></label>
<label>Answer type <select id="cc-type"><option value="">All</option>{opts}</select></label>
<label class="cc-check"><input id="cc-flagged" type="checkbox"> Only flagged for review ({n_flag})</label>
<p id="cc-count" class="meta" aria-live="polite"></p>
</div>
<div class="cc-scroll cc-tablewrap"><table class="cc-table cc-concepts{' cc-wide' if is_cdsi else ''}" id="cc-concepts"><thead><tr>{head}</tr></thead><tbody>
{chr(10).join(rows)}
</tbody></table></div>
{FILTER_JS}"""


def downloads_list(rel, cs):
    stem = f"CodeSystem-{cs['key'].replace('/', '-')}.json"
    items = [
        (stem, "FHIR CodeSystem (R4, JSON)", "The normalized terminology resource."),
        ("concepts.csv", "Source-derived CSV", "Every meaningful source column exactly as supplied, plus clearly labelled IVC columns where applicable."),
        ("provenance.json", "Provenance", "Source identity, checksums, retrieval date and generator version."),
        ("validation-report.html", "Validation report", "QA checks, anomalies, transformations and excluded spreadsheet structure."),
    ]
    return '<ul class="resource-list">' + "".join(f'<li><a href="{rel}{f}">{esc(t)}</a><br><span class="meta">{esc(d)}</span></li>' for f, t, d in items) + "</ul>"


def authority_page(cfg, cs, concepts, canonical, extra):
    ver = cs["version"]
    country = cfg["countries"][cs["country"]]
    tags = tag("draft", "Status: draft") + tag("poc", "Proof of concept") + tag("exp", "Experimental")
    meta = kv_table([
        ("Canonical URL", f"<code>{esc(canonical)}</code>"),
        ("Current version", f'{esc(ver)} <span class="meta">({esc(cs["versionNote"])})</span>'),
        ("Status", "Draft, proof of concept"),
        ("Publication date", esc(cfg["publicationDate"])),
        ("Concepts", f"{len(concepts)}"),
        ("Resource type", f'FHIR R4 ({esc(cfg["fhirVersion"])}) CodeSystem, content complete'),
        ("Primary language", "French, with English designations" if cs["language"] == "fr" else "English"),
    ])
    known = "".join(f"<li>{esc(k)}</li>" for k in cs["knownIssues"])
    labels = f"""<ul class="cc-legend">
<li>{tag('src', 'Source-supplied')} Codes and labels{', indication, contraindication and clarifying text, and related codes' if cs['source'] == 'cdsi' else ', and answer type'}.</li>
<li>{tag('ivc', 'IVC-added')} Canonical URL, version label, status, property definitions, FHIR structure{', and expected answer type' if cs['source'] == 'cdsi' else ''}. Table columns and values marked <span class="cc-mini">IVC</span> were supplied by IVC.</li>
<li>{tag('flag', 'IVC-flagged')} Notes in the table point out source content that needs review. The source content itself has not been changed.</li>
</ul>"""
    body = f"""<section class="section"><div class="wrap">
<div class="cc-tags">{tags}</div>
<div class="grid-2 cc-top">
<div><h2>About this code system</h2><p>{esc(cs['authorityDescription'])}</p><p>{esc(cs['description'])}</p>
<div class="callout"><p><strong>Operational assertions, not diagnoses.</strong> Each code asserts that a condition applies to a patient for immunization decision support. The sender decides whether the condition applies, and the receiving IIS or CDS engine uses the assertion. These codes are not problem-list entries, clinical findings or inference rules.</p></div></div>
<div class="card accent">{meta}</div>
</div></div></section>
<section class="section tint"><div class="wrap"><div class="section-head"><h2>Governance roles</h2><p>Under the Contextual Conditions proposal, these four roles can belong to different organizations.</p></div>{roles_table(cs['roles'])}</div></section>
<section class="section"><div class="wrap grid-2 cc-top">
<div><h2>Using these codes</h2>{hl7_examples(cs, canonical, concepts)}</div>
<div><h2>Downloads</h2><p>The current release ({esc(ver)}). The same files are kept unchanged under the <a href="{esc(ver)}/">{esc(ver)} version folder</a>, and a copy is always available under <a href="current/">current</a>.</p>{downloads_list(ver + '/', cs)}
<h3>Release history</h3><div class="cc-scroll"><table class="cc-table"><thead><tr><th scope="col">Version</th><th scope="col">Date</th><th scope="col">Status</th><th scope="col">Notes</th></tr></thead><tbody><tr><td><a href="{esc(ver)}/">{esc(ver)}</a></td><td>{esc(cfg['publicationDate'])}</td><td>Proof of concept</td><td>First IVC proof-of-concept publication.</td></tr></tbody></table></div>
{extra.get('history', '')}</div>
</div></section>
<section class="section tint"><div class="wrap"><div class="section-head"><h2>Concepts</h2><p>All {len(concepts)} concepts in source order. Each code has its own link, for example <code>{esc(canonical)}#c-{esc(concepts[0]['code'])}</code>.</p>{labels}</div>
{concept_table(cs, concepts)}</div></section>
<section class="section"><div class="wrap grid-2 cc-top">
<div><h2>Known issues</h2><ul>{known}</ul><p>See the <a href="{esc(ver)}/validation-report.html">validation report</a> for every automated check.</p>{extra.get('overview', '')}</div>
<div><h2>Copyright, licence and attribution</h2><p>{esc(cs['copyright'])}</p><h2>Corrections and feedback</h2><p>{esc(cs['feedback'])}</p><p>Contact IVC at <a href="mailto:{esc(cfg['contactEmail'])}?subject={esc('Contextual Conditions: ' + cs['title'])}">{esc(cfg['contactEmail'])}</a>.</p><p class="meta">Published codes are permanent. New codes may be added, but a published code is never reassigned to a different meaning.</p></div>
</div></section>"""
    crumbs = [("Contextual Conditions", "../../"), (country["name"], "../"), (cs["shortTitle"], None)]
    return page(cs["title"], cs["description"], 3, body, f"Contextual Conditions · {country['name']} · {cs['authority']}", cs["title"],
                f'Canonical code system <code class="cc-hero-code">{esc(canonical)}</code>', crumbs)


def version_page(cfg, cs, canonical, is_current):
    ver = cs["version"]
    country = cfg["countries"][cs["country"]]
    label = "current" if is_current else ver
    intro = (f"<p>The <code>current</code> folder holds a copy of the files for the current release, version {esc(ver)}. Link to the <a href=\"../{esc(ver)}/\">{esc(ver)}</a> folder if you need a reference that will never change.</p>"
             if is_current else f"<p>This folder holds version {esc(ver)}. Files in a published version folder are never changed. Later releases get their own folder.</p>")
    body = f"""<section class="section"><div class="wrap grid-2 cc-top">
<div>{intro}<p>To browse the concepts, return to the <a href="../">canonical page</a>.</p>
{kv_table([('Canonical URL', f'<code>{esc(canonical)}</code>'), ('Version', esc(ver)), ('Status', 'Draft, proof of concept'), ('Publication date', esc(cfg['publicationDate']))])}</div>
<div><h2>Files</h2>{downloads_list('', cs)}</div>
</div></section>"""
    crumbs = [("Contextual Conditions", "../../../"), (country["name"], "../../"), (cs["shortTitle"], "../"), (label, None)]
    return page(f"{cs['title']} {label}", f"{cs['title']} release files", 4, body, f"{cs['title']} · release files", f"Version {ver}" if not is_current else f"Current release ({ver})",
                "Machine-readable release artifacts", crumbs)


def validation_page(cfg, cs, report, canonical, depth_label):
    counts = report.counts()
    order = {"fail": 0, "warn": 1, "info": 2, "pass": 3}
    rows = "".join(f'<tr class="cc-v-{lvl}"><td>{tag(lvl, lvl.upper())}</td><td>{esc(chk)}</td><td>{esc(det)}</td></tr>'
                   for lvl, chk, det in sorted(report.items, key=lambda x: order[x[0]]))
    summary = " · ".join(f"{counts.get(k, 0)} {k}" for k in ("fail", "warn", "info", "pass"))
    body = f"""<section class="section"><div class="wrap">
<p>This report was generated automatically from the same normalized model as the CodeSystem, HTML and CSV. <strong>{esc(summary)}</strong>.</p>
<p class="meta">Levels: <b>fail</b> blocks a normative release. <b>warn</b> needs human review. <b>info</b> records a transformation or fact. <b>pass</b> means a check succeeded. Steps not yet automated: FHIR validator run (plan step 7) and named release approval (plan step 11).</p>
<div class="cc-scroll"><table class="cc-table cc-report"><thead><tr><th scope="col">Level</th><th scope="col">Check</th><th scope="col">Detail</th></tr></thead><tbody>{rows}</tbody></table></div>
<p><a href="./">Back to the release files</a> · <a href="../">Canonical page</a></p>
</div></section>"""
    country = cfg["countries"][cs["country"]]
    crumbs = [("Contextual Conditions", "../../../"), (country["name"], "../../"), (cs["shortTitle"], "../"), (depth_label, "./"), ("Validation report", None)]
    return page(f"{cs['title']} validation report", "Generated QA report", 4, body, f"{cs['title']} · {cs['version']}", "Validation report", esc(canonical), crumbs)


def country_page(cfg, code, systems, base):
    country = cfg["countries"][code]
    items = "".join(f'<li><a href="{esc(cs["authority"])}/">{esc(cs["title"])}</a><br><code>{esc(base + "/" + cs["key"])}</code><br><span class="meta">{esc(cs["roles"]["codeSystemAuthority"])}</span></li>' for cs in systems)
    future = country["illustrativeFutureAuthorities"]
    fut_html = ""
    if future:
        fut_html = "<h3>Room for more authorities</h3><p>Other organizations in this country could have their own code systems in this namespace. These examples are for illustration only and have <strong>not</strong> been assigned:</p><ul>" + "".join(
            f'<li><code>{esc(base)}/{code}/{esc(f["slug"])}</code>: {esc(f["name"])}</li>' for f in future) + "</ul>"
    body = f"""<section class="section"><div class="wrap grid-2 cc-top">
<div><h2>A coordination boundary</h2><p>The <code>{code}</code> segment comes from the ISO 3166-1 alpha-3 country code {esc(country['iso3166alpha3'])}. It is a delegated naming boundary. International coordination assigns country namespaces, and participants in each country coordinate their own authority identifiers.</p>
<p><strong>Being in this namespace does not mean a code system was authored or endorsed by the {esc(country['name'])} government.</strong> Each child code system names its own authority.</p>{fut_html}
<p class="meta">The process for coordinating names within a country is provisional (open decision 9).</p></div>
<div><h2>Code systems</h2><ul class="resource-list">{items}</ul></div>
</div></section>"""
    crumbs = [("Contextual Conditions", "../"), (country["name"], None)]
    return page(f"Contextual Conditions: {country['name']}", f"Contextual condition code systems in the {country['name']} namespace", 2, body, "Contextual Conditions · country namespace", country["name"],
                f'Namespace <code class="cc-hero-code">{esc(base)}/{code}</code>', crumbs)


def root_page(cfg, systems, base):
    cards = "".join(f'<div class="card accent{" green" if i % 2 else ""}"><p class="eyebrow">{esc(cfg["countries"][cs["country"]]["name"])}</p><h3><a href="{esc(cs["key"])}/">{esc(cs["title"])}</a></h3><p><code>{esc(base + "/" + cs["key"])}</code></p><p>{esc(cs["roles"]["codeSystemAuthority"])}. Version {esc(cs["version"])}, {cs["count"]} concepts.</p></div>' for i, cs in enumerate(systems))
    countries = "".join(f'<li><a href="{c}/">{esc(v["name"])}</a>: <code>{esc(base)}/{c}</code></li>' for c, v in cfg["countries"].items())
    props = "".join(f'<div id="{k}"><dt><code>{k}</code> <span class="meta">({t})</span></dt><dd>{esc(d)}</dd></div>' for k, (t, d) in PROPERTY_DEFS.items())
    decisions = "".join(f'<li{" class=cc-resolved" if d.get("resolved") else ""}><strong>{d["n"]}.</strong> {esc(d["text"])}{(" <em>" + esc(d["resolved"]) + "</em>") if d.get("resolved") else ""}</li>' for d in cfg["openDecisions"])
    body = f"""<section class="section"><div class="wrap grid-2 cc-top">
<div><h2>What is a contextual condition?</h2><p>A <strong>Contextual Condition (CC)</strong> is a standard coded assertion that a condition relevant to immunization applies to a patient, for use in clinical decision support. Examples include pregnancy, prior disease, immunocompromised status, travel risk, and whether a birth mother received an RSV vaccine during pregnancy.</p>
<p>Contextual conditions are operational inputs to CDS, not diagnoses, problem-list entries or inference rules. The sender decides whether a condition applies, and the receiving IIS or CDS engine uses the assertion. They are proposed as a third pillar of interoperability, next to <em>patient identification</em> and <em>immunization history</em>. New conditions can then be added as vocabulary changes, with no interface changes.</p></div>
<div class="card accent"><h3>Canonical namespace</h3><pre class="cc-code"><code>{esc(base)}
{esc(base)}/{{country}}/{{authority}}</code></pre><p>Each authority's code system gets a stable canonical URL. The URL identifies the code system inside exchanged data, and it also opens a page that explains the code system. Versioned files sit below it:</p><pre class="cc-code"><code>…/{{country}}/{{authority}}/            canonical page
…/{{country}}/{{authority}}/current/    current release files
…/{{country}}/{{authority}}/{{version}}/  release files that never change</code></pre></div>
</div></section>
<section class="section tint"><div class="wrap"><div class="section-head"><h2>Published code systems</h2><p>These two code systems describe the same general kind of concept, but they are independent. They are not coordinated, not equivalent, and not merged. Any future crosswalk will be published separately.</p></div><div class="grid-2">{cards}</div></div></section>
<section class="section"><div class="wrap grid-2 cc-top">
<div><h2>Governance roles</h2><p>Under the proposal, four roles can belong to different organizations:</p><dl class="cc-kv"><div><dt>Code System Authority</dt><dd>Defines what each concept means.</dd></div><div><dt>Publishing Authority</dt><dd>Publishes or distributes the vocabulary.</dd></div><div><dt>Namespace Authority</dt><dd>Maintains the canonical URLs and unique identifiers.</dd></div><div><dt>CDS Authority</dt><dd>Decides how concepts affect recommendation logic.</dd></div></dl>
<p>IVC coordinates the namespace and publication. It does not set clinical policy or national schedules, and it does not decide how CDS engines behave.</p></div>
<div><h2>Country namespaces</h2><ul>{countries}</ul><p>Country segments use ISO 3166-1 alpha-3 codes. A country segment is a coordination boundary. It does not mean a government authored or endorsed a code system.</p></div>
</div></section>
<section class="section tint"><div class="wrap"><div class="section-head"><h2>Concept properties</h2><p>Both code systems use these property definitions where they apply. IVC added <code>expectedValueType</code>, and its vocabulary is provisional.</p></div><dl class="definition-list cc-props">{props}</dl></div></section>
<section class="section"><div class="wrap narrow"><h2>Open decisions</h2><p>These questions are still open. This proof of concept exists to make them concrete and to invite review.</p><ol class="cc-decisions">{decisions}</ol>
<p>Send comments to <a href="mailto:{esc(cfg['contactEmail'])}?subject=Contextual%20Conditions%20proof%20of%20concept">{esc(cfg['contactEmail'])}</a>.</p></div></section>"""
    return page("Contextual Conditions", "Proof-of-concept canonical namespace for immunization contextual condition code systems", 1, body,
                "Proof of concept", "Contextual Conditions", "A canonical home for contextual condition code systems, the coded assertions that shape immunization decision support.")


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--stepintocdsi", type=Path, default=DEFAULT_STEPINTOCDSI)
    args = ap.parse_args()

    cfg = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    base = f"{cfg['siteBase']}/{cfg['rootPath']}"
    out_root = DOCS / cfg["rootPath"]
    us_rows, fr_rows = load_ivc_workbook()
    ivc_hash = sha256(IVC_WORKBOOK)

    built = []
    for cs in cfg["codeSystems"]:
        report = Report()
        canonical = f"{base}/{cs['key']}"
        if cs["source"] == "cdsi":
            cdsi = load_cdsi(args.stepintocdsi, report)
            if cdsi["release"] != cs["version"]:
                sys.exit(f"Config version {cs['version']} does not match StepIntoCDSI release {cdsi['release']}")
            concepts, csv_header, csv_rows = build_cdsi_model(cdsi, us_rows, report)
            provenance = {
                "source": "CDC Clinical Decision Support for Immunization (CDSi) Supporting Data",
                "sourceRelease": cdsi["release"],
                "sourceRetrievedAt": cdsi["manifest"].get("retrieved_at"),
                "sourceBundle": {"filename": cdsi["manifest"]["source_filename"], "sha256": cdsi["zipSha256"]},
                "sourceFilesUsed": [{"path": p, "sha256": h} for p, h in cdsi["checkedFiles"].items()],
                "normalizedBy": {"project": "StepIntoCDSI", "commit": cdsi["stepintoCommit"], "file": "normalized/schedules/schedule.json"},
                "ivcSupplied": {"file": "contextual-conditions/sources/ivc-supplied/Proposed Values.xlsx", "tab": "US", "sha256": ivc_hash, "fieldsUsed": ["Value (expected value type)"]},
            }
            latest = cdsi["history"][0] if cdsi["history"] else None
            extra = {}
            if latest:
                ch = "".join(f"<li><strong>{esc(c['area'])}</strong>: {esc(c['change'])} <span class=\"meta\">{esc(c['reason'])}</span></li>" for c in latest["changes"])
                extra["history"] = f"<h3>CDSi change history for {esc(latest['version'])}</h3><p class=\"meta\">From the official Coded Observations workbook, published {esc(latest['published'])}.</p><ul>{ch}</ul>"
            if cdsi["overview"]:
                extra["overview"] = "<h3>CDSi notes on these codes</h3><ul>" + "".join(f"<li>{esc(o)}</li>" for o in cdsi["overview"]) + "</ul><p class=\"meta\">Taken from the Overview tab of the official CDSi Coded Observations workbook.</p>"
        else:
            concepts, csv_header, csv_rows = build_syadem_model(fr_rows, report)
            provenance = {
                "source": "Syadem contextual conditions list, supplied to IVC by Syadem",
                "sourceRelease": None,
                "provisionalVersion": cs["version"],
                "ivcSupplied": {"file": "contextual-conditions/sources/ivc-supplied/Proposed Values.xlsx", "tab": "FR", "sha256": ivc_hash,
                                "fieldsUsed": ["Key", "English", "French", "Type"]},
                "permission": "Syadem authorized IVC to publish the complete list in both languages. The final licence and attribution wording is still to be confirmed.",
            }
            extra = {}

        cs["count"] = len(concepts)
        fhir = build_codesystem(cfg, cs, concepts, canonical, base)
        provenance.update({
            "codeSystem": canonical, "version": cs["version"], "publicationDate": cfg["publicationDate"], "status": "draft (proof of concept)",
            "generator": {"script": "contextual-conditions/generator/build.py", "version": GENERATOR_VERSION},
            "validationSummary": dict(report.counts()),
        })
        report.add("info", "Generated artifacts", "FHIR CodeSystem JSON, CSV, provenance JSON, HTML pages, this report")

        auth_dir = out_root / cs["key"]
        ver_dir = auth_dir / cs["version"]
        # Replace only this version and current/; earlier version folders are kept.
        for d in (ver_dir, auth_dir / "current"):
            if d.exists():
                shutil.rmtree(d)
        write_json(ver_dir / f"CodeSystem-{cs['key'].replace('/', '-')}.json", fhir)
        write_csv(ver_dir / "concepts.csv", csv_header, csv_rows)
        write_json(ver_dir / "provenance.json", provenance)
        write_text(ver_dir / "validation-report.html", validation_page(cfg, cs, report, canonical, cs["version"]))
        write_text(ver_dir / "index.html", version_page(cfg, cs, canonical, False))
        cur_dir = auth_dir / "current"
        shutil.copytree(ver_dir, cur_dir)
        write_text(cur_dir / "index.html", version_page(cfg, cs, canonical, True))
        write_text(cur_dir / "validation-report.html", validation_page(cfg, cs, report, canonical, "current"))
        write_text(auth_dir / "index.html", authority_page(cfg, cs, concepts, canonical, extra))
        c = report.counts()
        print(f"{cs['key']}: {len(concepts)} concepts; checks fail={c.get('fail', 0)} warn={c.get('warn', 0)} info={c.get('info', 0)} pass={c.get('pass', 0)}")
        built.append(cs)

    for code in cfg["countries"]:
        write_text(out_root / code / "index.html", country_page(cfg, code, [cs for cs in built if cs["country"] == code], base))
    write_text(out_root / "index.html", root_page(cfg, built, base))


if __name__ == "__main__":
    main()
